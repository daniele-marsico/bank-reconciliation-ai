"""Bank reconciliation: rules first, AI for what the rules can't solve.

Pass 1  Reference   - invoice number (or its numeric part) in the bank description
Pass 2  Name+amount - same amount, counterparty name similar, date close to due date
Pass 3  Combined    - one transfer = sum of two open invoices of the same counterparty
Pass 4  AI          - remaining lines are sent to Claude, which proposes matches
                      with a confidence score and a reason (only if ANTHROPIC_API_KEY
                      is set). AI proposals are marked "to review", never auto-booked.

Usage:
    python reconcile.py
    python reconcile.py --fee-tolerance 20 --no-ai
"""
from __future__ import annotations

import argparse
import json
import os
import re
from dataclasses import dataclass, field
from datetime import date
from difflib import SequenceMatcher
from itertools import combinations
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from io_utils import find_file, read_any, standardise, to_date, to_number

BASE = Path(__file__).parent
DATA_DIR = BASE / "data"
OUTPUT_DIR = BASE / "output"

DATE_WINDOW_DAYS = 15
NAME_SIMILARITY = 0.55
LEGAL_FORMS = r"\b(SRL|SPA|SNC|SAS|GMBH|AB|SA|LTD|&|FIGLI)\b"


@dataclass
class Match:
    line: int
    documents: list[str]
    rule: str
    difference: float = 0.0
    confidence: float = 1.0
    note: str = ""


@dataclass
class State:
    bank: pd.DataFrame
    items: pd.DataFrame
    matches: list[Match] = field(default_factory=list)

    @property
    def open_lines(self) -> pd.DataFrame:
        done = {m.line for m in self.matches}
        return self.bank[~self.bank["Line"].isin(done)]

    @property
    def open_items(self) -> pd.DataFrame:
        done = {d for m in self.matches if m.rule != "AI suggestion" for d in m.documents}
        return self.items[~self.items["Document No."].isin(done)]


# ----------------------------------------------------------------- data ----
# Column names accepted for each field (home-banking and Business Central, EN and IT)
BANK_ALIASES = {
    "Booking Date": ["Data contabile", "Data operazione", "Data", "Date"],
    "Description": ["Descrizione", "Causale", "Descrizione operazione", "Dettagli"],
    "Amount": ["Importo"],
    "Credit": ["Entrate", "Avere", "Accrediti"],
    "Debit": ["Uscite", "Dare", "Addebiti"],
}
ITEM_ALIASES = {
    "Document No.": ["Nr. documento", "N. documento", "Numero documento"],
    "Counterparty": ["Customer Name", "Vendor Name", "Nome cliente", "Nome fornitore", "Nome", "Ragione sociale"],
    "Due Date": ["Data scadenza", "Scadenza"],
    "Amount": ["Remaining Amount", "Importo residuo", "Importo"],
    "Type": ["Tipo documento", "Document Type"],
}


def load_data() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Bank statement and open items from data/ (CSV or Excel, EN or IT columns)."""
    bank = standardise(read_any(find_file(DATA_DIR, "bank_statement")), BANK_ALIASES,
                       ["Booking Date", "Description"])
    if "Amount" not in bank.columns:  # separate in / out columns, as many banks export
        if not {"Credit", "Debit"} <= set(bank.columns):
            raise ValueError("Bank file needs an Amount column, or two columns for money in and out")
        bank["Amount"] = to_number(bank["Credit"]).fillna(0) - to_number(bank["Debit"]).fillna(0).abs()
    bank["Amount"] = to_number(bank["Amount"]).round(2)
    bank["Booking Date"] = to_date(bank["Booking Date"]).dt.date
    bank["Description"] = bank["Description"].fillna("").astype(str)
    bank = bank.dropna(subset=["Amount"])
    if "Line" not in bank.columns:
        bank.insert(0, "Line", range(1, len(bank) + 1))
    bank["Line"] = bank["Line"].astype(int)
    bank = bank[["Line", "Booking Date", "Description", "Amount"]]

    items = standardise(read_any(find_file(DATA_DIR, "open_items")), ITEM_ALIASES,
                        ["Document No.", "Counterparty", "Due Date", "Amount"])
    items["Amount"] = to_number(items["Amount"]).round(2)
    items["Due Date"] = to_date(items["Due Date"]).dt.date
    items["Document No."] = items["Document No."].astype(str).str.strip()
    if "Type" not in items.columns:
        items["Type"] = items["Amount"].map(lambda v: "Customer invoice" if v >= 0 else "Vendor invoice")
    items = items[["Document No.", "Type", "Counterparty", "Due Date", "Amount"]]
    return bank.reset_index(drop=True), items.dropna(subset=["Amount"]).reset_index(drop=True)


# -------------------------------------------------------------- helpers ----
def clean_name(text: str) -> str:
    text = re.sub(LEGAL_FORMS, " ", text.upper())
    return " ".join(re.sub(r"[^A-Z ]", " ", text).split())


def name_score(counterparty: str, description: str) -> float:
    """How well the counterparty name appears inside the bank description (0-1)."""
    name, desc = clean_name(counterparty), clean_name(description)
    if not name:
        return 0.0
    if name in desc:
        return 1.0
    # best similarity of the name against any same-length window of the description
    words, n = desc.split(), len(name.split())
    windows = [" ".join(words[i:i + n]) for i in range(max(1, len(words) - n + 1))]
    return max(SequenceMatcher(None, name, w).ratio() for w in windows)


def days_apart(a, b) -> int:
    return abs((pd.Timestamp(a) - pd.Timestamp(b)).days)


# --------------------------------------------------------------- passes ----
def pass_reference(s: State, fee_tol: float) -> None:
    for _, b in s.open_lines.iterrows():
        desc = b["Description"].upper()
        for _, it in s.open_items.iterrows():
            doc = it["Document No."]
            num = re.sub(r"\D", "", doc.split("-")[-1])
            if doc.upper() in desc or re.search(rf"\b{num}\b", desc):
                diff = round(b["Amount"] - it["Amount"], 2)
                if diff == 0:
                    s.matches.append(Match(b["Line"], [doc], "Reference"))
                elif -fee_tol <= diff < 0:
                    s.matches.append(Match(b["Line"], [doc], "Reference", diff, note="Bank fee deducted"))
                else:
                    continue
                break


def pass_name_amount(s: State, fee_tol: float) -> None:
    for _, b in s.open_lines.iterrows():
        best = None
        for _, it in s.open_items.iterrows():
            diff = round(b["Amount"] - it["Amount"], 2)
            if not (diff == 0 or -fee_tol <= diff < 0):
                continue
            if days_apart(b["Booking Date"], it["Due Date"]) > DATE_WINDOW_DAYS:
                continue
            score = name_score(it["Counterparty"], b["Description"])
            if score >= NAME_SIMILARITY and (best is None or score > best[0]):
                best = (score, it["Document No."], diff)
        if best:
            note = "Bank fee deducted" if best[2] else ""
            s.matches.append(Match(b["Line"], [best[1]], "Name + amount", best[2], round(best[0], 2), note))


def pass_combined(s: State) -> None:
    for _, b in s.open_lines.iterrows():
        items = s.open_items
        for (_, i1), (_, i2) in combinations(items.iterrows(), 2):
            if i1["Counterparty"] != i2["Counterparty"]:
                continue
            if round(i1["Amount"] + i2["Amount"], 2) != b["Amount"]:
                continue
            if name_score(i1["Counterparty"], b["Description"]) < NAME_SIMILARITY:
                continue
            s.matches.append(Match(b["Line"], [i1["Document No."], i2["Document No."]], "Combined payment"))
            break


def pass_ai(s: State, model: str) -> None:
    lines, items = s.open_lines, s.open_items
    if lines.empty:
        return
    try:
        import anthropic
    except ImportError:
        print("anthropic package not installed - skipping AI pass")
        return

    prompt = f"""You are helping an accountant reconcile a bank statement.

Unmatched bank lines (positive = incoming):
{lines.to_csv(index=False)}
Open ledger items (customer invoices positive, vendor invoices negative):
{items.to_csv(index=False)}
For every bank line propose the most likely ledger document(s), or none.
Typical cases: partial payments, bank fees, combined payments, bank charges
(no document, suggest the G/L account category instead), unknown payers.

Reply with JSON only, no prose, in this shape:
[{{"line": 1, "documents": ["INV-..."], "confidence": 0.0-1.0, "reason": "short explanation"}}]"""

    client = anthropic.Anthropic()
    resp = client.messages.create(model=model, max_tokens=2000,
                                  messages=[{"role": "user", "content": prompt}])
    text = "".join(block.text for block in resp.content if block.type == "text")
    for p in parse_ai_reply(text, set(lines["Line"]), set(items["Document No."])):
        s.matches.append(Match(p["line"], p["documents"], "AI suggestion",
                               confidence=p["confidence"], note=p["reason"]))


def parse_ai_reply(text: str, valid_lines: set, valid_docs: set) -> list[dict]:
    """Extract and validate the JSON list returned by the model."""
    m = re.search(r"\[.*\]", text, re.S)
    if not m:
        return []
    out = []
    for p in json.loads(m.group(0)):
        if p.get("line") not in valid_lines:
            continue
        docs = [d for d in p.get("documents") or [] if d in valid_docs]
        out.append({"line": p["line"], "documents": docs,
                    "confidence": float(p.get("confidence", 0)), "reason": str(p.get("reason", ""))})
    return out


# --------------------------------------------------------------- output ----
def write_excel(s: State) -> Path:
    by_line = {m.line: m for m in s.matches}
    rows = []
    for _, b in s.bank.iterrows():
        m = by_line.get(b["Line"])
        if m is None:
            status = "Unmatched - to review"
        elif m.rule == "AI suggestion":
            status = "AI suggestion - to review"
        else:
            status = "Matched"
        rows.append({**b.to_dict(), "Status": status, "Rule": m.rule if m else "",
                     "Documents": ", ".join(m.documents) if m else "",
                     "Difference": m.difference if m else None,
                     "Confidence": m.confidence if m else None, "Note": m.note if m else ""})
    result = pd.DataFrame(rows)
    open_items = s.open_items

    summary = pd.DataFrame([
        {"Metric": "Bank lines", "Value": len(result)},
        {"Metric": "Matched automatically", "Value": int((result["Status"] == "Matched").sum())},
        {"Metric": "AI suggestions to review", "Value": int((result["Status"].str.startswith("AI")).sum())},
        {"Metric": "Unmatched", "Value": int((result["Status"].str.startswith("Unmatched")).sum())},
        {"Metric": "Open items still unpaid", "Value": len(open_items)},
        {"Metric": "Bank fees found", "Value": round(result["Difference"].fillna(0).sum(), 2)},
    ])
    for rule, n in result[result["Rule"] != ""]["Rule"].value_counts().items():
        summary.loc[len(summary)] = {"Metric": f"  by rule: {rule}", "Value": int(n)}

    fills = {"Matched": "C6EFCE", "AI suggestion - to review": "FFEB9C", "Unmatched - to review": "F8CBAD"}
    wb = Workbook()
    for title, df in [("Summary", summary), ("Bank lines", result), ("Open items", open_items)]:
        ws = wb.active if title == "Summary" else wb.create_sheet(title)
        ws.title = title
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(horizontal="center")
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
            if title == "Bank lines":
                status = row[list(df.columns).index("Status")]
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor=fills[status])
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 60)
            for c in col[1:]:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
        ws.freeze_panes = "A2"

    OUTPUT_DIR.mkdir(exist_ok=True)
    out = OUTPUT_DIR / f"reconciliation_{date.today().isoformat()}.xlsx"
    wb.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--fee-tolerance", type=float, default=20.0,
                        help="Max bank fee accepted as a deduction (default 20)")
    parser.add_argument("--no-ai", action="store_true", help="Skip the AI pass")
    parser.add_argument("--model", default=os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-5"))
    args = parser.parse_args()

    if not any(DATA_DIR.glob("bank_statement.*")):
        import generate_sample_data
        generate_sample_data.main()

    bank, items = load_data()
    s = State(bank, items)

    pass_reference(s, args.fee_tolerance)
    pass_name_amount(s, args.fee_tolerance)
    pass_combined(s)
    if not args.no_ai and os.getenv("ANTHROPIC_API_KEY"):
        pass_ai(s, args.model)
    elif not args.no_ai:
        print("ANTHROPIC_API_KEY not set - AI pass skipped (rules only)")

    out = write_excel(s)
    auto = sum(1 for m in s.matches if m.rule != "AI suggestion")
    print(f"{auto}/{len(bank)} bank lines matched by rules -> {out}")


if __name__ == "__main__":
    main()
